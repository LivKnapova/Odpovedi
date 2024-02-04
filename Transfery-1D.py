#!/usr/bin/env python
# coding: utf-8

# In[33]:


import pandas as pd
import docx


# In[7]:


# open file
df = pd.read_csv("C:/Users/Admin/OneDrive/Desktop/transfery.csv")

df.head()


# In[5]:


df['genetic_method'] = df['genetic_method'].astype("string")


# In[29]:


count_of_PGTA = df[df['genetic_method'] == 'PGT-A'].count()
print(count_of_PGTA['genetic_method'])
count_of_karyomap = df[df['genetic_method'] == 'Karyomapping'].count()
print(count_of_karyomap['genetic_method'])
count_of_NaN = df['genetic_method'].isna().sum()
print(count_of_NaN)
count_of_PGTSR = df[df['genetic_method'] == 'PGT-SR'].count()
print(count_of_PGTSR['genetic_method'])
count_of_OneGene = df[df['genetic_method'] == 'OneGene'].count()
print(count_of_OneGene['genetic_method'])
count_of_others = df[(df['genetic_method'] != 'PGT-A') & (df['genetic_method'] != 'Karyomapping')& (df['genetic_method'] != 'PGT-SR')& (df['genetic_method'] != 'OneGene')& (df['genetic_method'] != 'NaN')].count()
print(count_of_others['genetic_method'])


# In[30]:


df4 = pd.DataFrame(columns=['PGT-A', 'PGT-SR', 'Karyomapping', 'OneGene', 'Without value', 'Others' ])
df4['PGT-A'] = [count_of_PGTA['genetic_method']]
df4['PGT-SR'] = [count_of_PGTSR['genetic_method']]
df4['Karyomapping'] = [count_of_karyomap['genetic_method']]
df4['OneGene'] = [count_of_OneGene['genetic_method']]
df4['Without value'] = [count_of_NaN]
df4['Others'] = [count_of_others['genetic_method']]
df4


# In[38]:


df4.to_csv('3Results.csv', sep=',', index=False, encoding='utf-8')


# In[34]:


doc = docx.Document('C:/Users/Admin/OneDrive/Desktop/Results.docx')


# In[35]:


t = doc.add_table(df4.shape[0]+1, df4.shape[1])

for j in range(df4.shape[-1]):
    t.cell(0,j).text = df4.columns[j]
    
for i in range(df4.shape[0]):
    for j in range(df4.shape[-1]):
        t.cell(i+1,j).text = str(df4.values[i,j])


# In[36]:


doc.save('C:/Users/Admin/OneDrive/Desktop/Results.docx')


# In[ ]:




